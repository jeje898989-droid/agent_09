"""Reusable helpers for building .pptx decks with python-pptx (mk-ppt skill).

Import from a deck script:
    import sys; sys.path.insert(0, ".claude/skills/mk-ppt/scripts")
    from pptx_kit import *

CLI:
    python .claude/skills/mk-ppt/scripts/pptx_kit.py next-version presentation AI_발전_발표
    python .claude/skills/mk-ppt/scripts/pptx_kit.py inspect presentation/AI_발전_발표_v1.pptx
"""
import copy
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Windows consoles default to cp949, which garbles Korean output and paths passed through pipes
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")

FONT = "맑은 고딕"
THEME = {
    "primary": RGBColor(0x1F, 0x38, 0x64),
    "accent": RGBColor(0x2E, 0x75, 0xB6),
    "text": RGBColor(0x26, 0x26, 0x26),
    "muted": RGBColor(0x7F, 0x7F, 0x7F),
    "light": RGBColor(0xF2, 0xF2, 0xF2),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

# Layout indexes of the default python-pptx template
LAYOUT_TITLE, LAYOUT_TITLE_CONTENT, LAYOUT_SECTION, LAYOUT_TITLE_ONLY, LAYOUT_BLANK = 0, 1, 2, 5, 6


# ---------- files & versions ----------

def next_version_path(folder, base, ext=".pptx"):
    """Return folder/base_vN.ext with N one higher than any existing version."""
    os.makedirs(folder, exist_ok=True)
    pat = re.compile(re.escape(base) + r"_v(\d+)" + re.escape(ext) + "$")
    nums = [int(m.group(1)) for f in os.listdir(folder) if (m := pat.match(f))]
    return os.path.join(folder, f"{base}_v{max(nums, default=0) + 1}{ext}")


def save(prs, path):
    """Save, refusing to overwrite an existing version."""
    if os.path.exists(path):
        raise FileExistsError(f"{path} exists; use next_version_path() instead")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    prs.save(path)
    return path


def new_deck(widescreen=True, template=None):
    """New presentation; 16:9 by default, or based on a template .pptx/.potx."""
    prs = Presentation(template) if template else Presentation()
    if widescreen and not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


# ---------- text ----------

def style_run(run, size=None, bold=None, color=None, font=FONT, italic=None):
    f = run.font
    f.name = font
    # East Asian font must be set separately or Korean falls back to the theme font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {})
        rPr.append(ea)
    ea.set("typeface", font)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def set_text(text_frame, text, size=18, bold=False, color=None, align=None):
    """Replace a text frame's content with one styled paragraph."""
    text_frame.clear()
    p = text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    style_run(run, size, bold, color or THEME["text"])
    if align is not None:
        p.alignment = align
    return p


def add_paragraphs(text_frame, items, size=18, color=None, space_after=6):
    """Fill a text frame with bullet items. An item is "text" or ("text", level)."""
    text_frame.clear()
    text_frame.word_wrap = True
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.level = level
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        style_run(run, size - 2 * level, color=color or THEME["text"])


def add_textbox(slide, left, top, width, height, text, size=16, bold=False, color=None, align=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.text_frame.word_wrap = True
    set_text(tb.text_frame, text, size, bold, color, align)
    return tb


def _title(slide, text, size=32):
    set_text(slide.shapes.title.text_frame, text, size, True, THEME["primary"])


# ---------- slide builders ----------

def add_title_slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE])
    set_text(s.shapes.title.text_frame, title, 40, True, THEME["primary"], PP_ALIGN.CENTER)
    sub = s.placeholders[1]
    if subtitle:
        set_text(sub.text_frame, subtitle, 20, color=THEME["muted"], align=PP_ALIGN.CENTER)
    else:
        sub._element.getparent().remove(sub._element)
    return s


def add_section_slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_SECTION])
    _title(s, title, 36)
    if len(s.placeholders) > 1:
        ph = s.placeholders[1]
        if subtitle:
            set_text(ph.text_frame, subtitle, 18, color=THEME["muted"])
        else:
            ph._element.getparent().remove(ph._element)
    return s


def add_bullet_slide(prs, title, items, size=20):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_CONTENT])
    _title(s, title)
    body = s.placeholders[1]
    body.left, body.top = Inches(0.7), Inches(1.6)
    body.width, body.height = prs.slide_width - Inches(1.4), prs.slide_height - Inches(2.2)
    add_paragraphs(body.text_frame, items, size)
    return s


def _content_box(prs):
    """left, top, width, height of the area under the title."""
    return Inches(0.7), Inches(1.6), prs.slide_width - Inches(1.4), prs.slide_height - Inches(2.2)


def add_table_slide(prs, title, rows, col_widths=None, font_size=14, header=True):
    """rows: list of lists of strings; first row is the header."""
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _title(s, title)
    left, top, width, _ = _content_box(prs)
    n_rows, n_cols = len(rows), len(rows[0])
    gf = s.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.45) * n_rows)
    table = gf.table
    table.first_row = header
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            table.columns[i].width = Emu(int(width * w / total))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and r == 0
            set_text(cell.text_frame, str(val), font_size, is_head,
                     THEME["white"] if is_head else THEME["text"])
            if is_head:
                cell.fill.solid()
                cell.fill.fore_color.rgb = THEME["primary"]
    return s, table


CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "radar": XL_CHART_TYPE.RADAR,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


def add_chart_slide(prs, title, categories, series, kind="column", number_format=None,
                    legend=True, data_labels=False):
    """series: {"name": [values...]}. For scatter, series values are [(x, y), ...] and categories is ignored."""
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _title(s, title)
    if kind == "scatter":
        data = XyChartData()
        for name, points in series.items():
            ser = data.add_series(name)
            for x, y in points:
                ser.add_data_point(x, y)
    else:
        data = CategoryChartData()
        data.categories = categories
        for name, values in series.items():
            data.add_series(name, values)
    chart = s.shapes.add_chart(CHART_TYPES.get(kind, kind), *_content_box(prs), data).chart
    chart.font.name, chart.font.size = FONT, Pt(14)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if data_labels:
        plot = chart.plots[0]
        plot.has_data_labels = True
        if number_format:
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
        if kind in ("pie", "doughnut"):
            plot.data_labels.show_percentage = number_format is None
    if number_format and kind not in ("pie", "doughnut", "scatter", "radar"):
        chart.value_axis.tick_labels.number_format = number_format
        chart.value_axis.tick_labels.number_format_is_linked = False
    return s, chart


def add_image_slide(prs, title, image_path, caption=None):
    """Picture scaled to fit the content area, centered."""
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _title(s, title)
    left, top, width, height = _content_box(prs)
    if caption:
        height -= Inches(0.5)
    pic = s.shapes.add_picture(image_path, left, top)
    ratio = min(width / pic.width, height / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int(left + (width - pic.width) / 2)
    pic.top = int(top + (height - pic.height) / 2)
    if caption:
        add_textbox(s, left, top + height + Inches(0.1), width, Inches(0.4), caption, 14,
                    color=THEME["muted"], align=PP_ALIGN.CENTER)
    return s, pic


def add_two_column_slide(prs, title, left_items, right_items, left_head=None, right_head=None):
    s = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    _title(s, title)
    left, top, width, height = _content_box(prs)
    gap = Inches(0.4)
    col_w = int((width - gap) / 2)
    for i, (items, head) in enumerate(((left_items, left_head), (right_items, right_head))):
        x = left + i * (col_w + gap)
        y = top
        if head:
            box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, col_w, Inches(0.5))
            box.fill.solid()
            box.fill.fore_color.rgb = THEME["accent"]
            box.line.fill.background()
            set_text(box.text_frame, head, 18, True, THEME["white"], PP_ALIGN.CENTER)
            y += Inches(0.6)
        tb = s.shapes.add_textbox(x, y, col_w, top + height - y)
        add_paragraphs(tb.text_frame, items, 18)
    return s


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_background(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_footer_numbers(prs, skip_first=True, size=10):
    """Write 'n / total' at the bottom-right of every slide."""
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        if skip_first and i == 1:
            continue
        add_textbox(slide, prs.slide_width - Inches(1.6), prs.slide_height - Inches(0.5),
                    Inches(1.3), Inches(0.3), f"{i} / {total}", size,
                    color=THEME["muted"], align=PP_ALIGN.RIGHT)


# ---------- slide management (not in the public API; XML workarounds) ----------

def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


def move_slide(prs, old_index, new_index):
    sldIdLst = prs.slides._sldIdLst
    el = sldIdLst[old_index]
    sldIdLst.remove(el)
    sldIdLst.insert(new_index, el)


def duplicate_slide(prs, index):
    """Append a copy of slide[index]. Copies shapes and picture/media relationships;
    charts are shared with the source, so call replace_data on a fresh chart instead if needed."""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dst.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    for rel in src.part.rels.values():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            dst.part.rels._rels[rel.rId] = rel  # keep same rId so copied XML still resolves
    return dst


# ---------- verification ----------

def inspect(path):
    """Print an outline of a deck: per slide, each shape's kind and text."""
    prs = Presentation(path)
    print(f"{path}: {len(prs.slides)} slides, "
          f"{prs.slide_width / 914400:.2f}x{prs.slide_height / 914400:.2f} in")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n[{i}] layout={slide.slide_layout.name}")
        for shp in slide.shapes:
            kind = ("chart" if shp.has_chart else "table" if shp.has_table
                    else "picture" if shp.shape_type == 13 else "shape")
            text = shp.text_frame.text.replace("\n", " | ")[:100] if shp.has_text_frame else ""
            if shp.has_table:
                t = shp.table
                text = f"{len(t.rows)}x{len(t.columns)}: " + " | ".join(c.text for c in t.rows[0].cells)
            if shp.has_chart:
                ch = shp.chart
                text = f"{ch.chart_type} series={[s.name for s in ch.plots[0].series]}"
            off = (shp.left + shp.width > prs.slide_width or shp.top + shp.height > prs.slide_height)
            print(f"  - {kind:8} {shp.name}: {text}{'  !! OFF-SLIDE' if off else ''}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
            print(f"  notes: {slide.notes_slide.notes_text_frame.text[:80]}")


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "inspect":
        inspect(sys.argv[2])
    elif len(sys.argv) >= 4 and sys.argv[1] == "next-version":
        print(next_version_path(sys.argv[2], sys.argv[3]))
    else:
        print(__doc__)
