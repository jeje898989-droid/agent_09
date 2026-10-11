"""AI 발전 동향 발표자료 생성 스크립트 (8장).

디자인: references/ppt/slide-design-system.md (흰 배경, 4색 액센트 카드 레이아웃)
규칙:   .claude/rules/ppt-rule-v1/ppt-rule-v1.md
출처:   research/ai-development-research-v2.md

실행: python scripts/make_ppt_ai_development.py
"""
import math
import sys
import unicodedata

sys.path.insert(0, ".claude/skills/mk-ppt/scripts")
from pptx_kit import new_deck, next_version_path, save, set_notes, LAYOUT_BLANK

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------- 디자인 토큰 (slide-design-system.md) ----------
YELLOW = RGBColor(0xFB, 0xB9, 0x2F)
CORAL = RGBColor(0xFB, 0x85, 0x67)
PINK = RGBColor(0xFB, 0x56, 0x9B)
PURPLE = RGBColor(0x9C, 0x28, 0xB1)
ACCENTS = [YELLOW, CORAL, PINK, PURPLE]
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 디자인 파일: 한글 슬라이드는 Noto Sans KR. Fira Sans Extra Condensed가 없는 PC에서는 대체 글꼴이
# 영문 자간을 벌려 읽기 어려우므로, 제목류도 Noto Sans KR로 통일한다.
HEAD_FONT = "Noto Sans KR"
BODY_FONT = "Roboto"
KO_FONT = "Noto Sans KR"

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.4                       # ppt-rule-v1: 가장자리 최소 여백
CARD_X0, CARD_STEP, CARD_W = 1.03, 2.89, 2.62
HEAD_TOP, HEAD_H = 1.9, 0.85
CARD_TOP, CARD_H = 2.93, 3.52


# ---------- 기본 도구 ----------

def style(run, size, bold=False, color=BLACK, latin=BODY_FONT):
    f = run.font
    f.name, f.size, f.bold = latin, Pt(size), bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", KO_FONT)


def text(slide, x, y, w, h, lines, size, bold=False, color=BLACK, latin=BODY_FONT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, name=None, space_after=0):
    """텍스트 상자. lines는 문자열 또는 문자열 리스트(한 항목 = 한 문단)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        style(p.add_run(), size, bold, color, latin)
        p.runs[0].text = line
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.5, name=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    return shp


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[LAYOUT_BLANK])


def title(slide, t):
    text(slide, MARGIN + 0.2, 0.7, SLIDE_W - 2 * (MARGIN + 0.2), 0.8, t, 36, True,
         latin=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE, name="Title")


def page_no(slide, n, total):
    text(slide, SLIDE_W - 1.9, 6.75, 1.2, 0.3, f"{n} / {total}", 14,
         align=PP_ALIGN.RIGHT, name="PageNo")


def dots(slide, y):
    """표지·마무리 장식: 액센트 4색 막대."""
    w, gap = 0.6, 0.15
    x0 = (SLIDE_W - (4 * w + 3 * gap)) / 2
    for i, c in enumerate(ACCENTS):
        bar = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x0 + i * (w + gap), y, w, 0.14,
                    fill=c, name=f"Deco{i + 1}")
        bar.adjustments[0] = 0.5


def icon(slide, kind, cx, y, color, label=None, w=0.7, h=0.7):
    """라인 아이콘: 채우기 없이 액센트색 2pt 외곽선. label이 있으면 원 안에 숫자."""
    ic = shape(slide, kind, cx - w / 2, y + (0.7 - h) / 2, w, h, line=color, line_w=2, name="Icon")
    if label:
        tf = ic.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        style(p.add_run(), 22, True, color, HEAD_FONT)
        p.runs[0].text = label


def cards(slide, items):
    """시그니처 레이아웃: 헤더 막대 + 흰 카드(아이콘·부제·본문). 2~4개는 가운데 정렬."""
    n = len(items)
    offset = (4 - n) * CARD_STEP / 2
    for i, it in enumerate(items):
        x = CARD_X0 + offset + i * CARD_STEP
        cx = x + CARD_W / 2
        color = ACCENTS[i]
        head = shape(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, HEAD_TOP, CARD_W, HEAD_H,
                     fill=color, name=f"Header{i + 1}")
        head.adjustments[0] = 0.25
        text(slide, x, HEAD_TOP, CARD_W, HEAD_H, it["head"], 22, True, WHITE, HEAD_FONT,
             anchor=MSO_ANCHOR.MIDDLE, name=f"HeaderText{i + 1}")
        card = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, CARD_TOP, CARD_W, CARD_H,
                     fill=WHITE, line=BLACK, name=f"Card{i + 1}")
        card.adjustments[0] = 0.1 / CARD_W
        icon(slide, it["icon"], cx, CARD_TOP + 0.35, color, it.get("label"),
             *it.get("icon_size", (0.7, 0.7)))
        text(slide, x + 0.05, CARD_TOP + 1.25, CARD_W - 0.1, 0.55, it["sub"], 22, True,
             color, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE, name=f"Sub{i + 1}")
        text(slide, x + 0.05, CARD_TOP + 1.9, CARD_W - 0.1, 1.4, it["body"], 16,
             name=f"Body{i + 1}")


def table_borders(cell, color="000000", w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for idx, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.find(qn(tag))
        if ln is not None:
            tcPr.remove(ln)
        ln = tcPr.makeelement(qn(tag), {"w": str(int(w_pt * 12700)), "cmpd": "sng"})
        solid = ln.makeelement(qn("a:solidFill"), {})
        clr = solid.makeelement(qn("a:srgbClr"), {"val": color})
        solid.append(clr)
        ln.append(solid)
        tcPr.insert(idx, ln)  # 스키마 순서: 선(lnL…lnB)이 채우기보다 앞에 와야 표시된다


# ---------- 슬라이드 ----------
TOTAL = 8
prs = new_deck()
SRC = "research/ai-development-research-v2.md"

# 1. 표지
s = blank(prs)
text(s, 1.0, 2.2, SLIDE_W - 2.0, 1.2, "AI 발전 동향", 54, True, latin=HEAD_FONT,
     anchor=MSO_ANCHOR.MIDDLE, name="CoverTitle")
text(s, 1.0, 3.45, SLIDE_W - 2.0, 0.7, "2024–2026 기술·산업·규제 핵심 정리", 26, True,
     YELLOW, HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE, name="CoverSub")
dots(s, 4.5)
text(s, 1.0, 5.6, SLIDE_W - 2.0, 0.45, f"2026년 10월 · 출처: {SRC}", 14, name="CoverSource")
set_notes(s, "오늘은 2024년부터 2026년까지의 AI 발전 흐름을 기술, 산업·경제, 규제·윤리, "
             "전망의 네 갈래로 정리해 말씀드리겠습니다. 수치는 Stanford AI Index 2026 등 "
             "공개 자료를 기반으로 한 자료조사 노트에서 가져왔습니다.")

# 2. 발표 순서
s = blank(prs)
title(s, "발표 순서")
cards(s, [
    dict(head="PART 1", icon=MSO_SHAPE.OVAL, label="1", sub="핵심 기술",
         body=["모델 성능", "에이전트·미중 경쟁"]),
    dict(head="PART 2", icon=MSO_SHAPE.OVAL, label="2", sub="산업·경제",
         body=["투자·도입", "생산성·노동시장"]),
    dict(head="PART 3", icon=MSO_SHAPE.OVAL, label="3", sub="규제·윤리",
         body=["EU AI Act", "한국 AI 기본법"]),
    dict(head="PART 4", icon=MSO_SHAPE.OVAL, label="4", sub="전망",
         body=["4대 핵심 이슈", "시사점"]),
])
page_no(s, 2, TOTAL)
set_notes(s, "발표는 네 파트로 진행합니다. 기술 동향을 먼저 보고, 그 기술이 투자와 일자리에 "
             "어떤 영향을 주는지, 각국이 어떻게 규제하고 있는지를 본 뒤 전망으로 마무리합니다.")

# 3. 핵심 기술 트렌드
s = blank(prs)
title(s, "핵심 기술 트렌드")
cards(s, [
    dict(head="모델 성능", icon=MSO_SHAPE.STAR_5_POINT, sub="코딩 ~100%",
         body=["SWE-bench 점수", "1년 새 60%→~100%"]),
    dict(head="AI 에이전트", icon=MSO_SHAPE.GEAR_6, sub="성공률 66%",
         body=["OSWorld 과제", "12% → 66%로 상승"]),
    dict(head="미·중 경쟁", icon=MSO_SHAPE.LEFT_RIGHT_ARROW, icon_size=(1.0, 0.5), sub="격차 2.7%",
         body=["성능은 사실상 대등", "투자는 미국이 23배"]),
    dict(head="성능 편차", icon=MSO_SHAPE.LIGHTNING_BOLT, sub="시계 50.1%",
         body=["IMO 금메달 모델도", "시계 읽기는 절반만"]),
])
page_no(s, 3, TOTAL)
set_notes(s, "2025년 주목할 프런티어 모델의 90% 이상을 기업이 만들었고, 박사급 과학 문제와 "
             "수학 경시에서 인간 기준을 넘어섰습니다. 컴퓨터 사용 에이전트의 과제 성공률도 "
             "12%에서 66%로 뛰었습니다. 미·중 성능 격차는 2026년 3월 기준 2.7%로 사실상 "
             "사라졌지만, 민간 투자는 미국이 2,859억 달러로 중국의 23배입니다. 다만 IMO "
             "금메달 모델이 아날로그 시계는 50.1%만 맞히는 '들쭉날쭉한 프런티어'가 있어, "
             "벤치마크 점수만으로 실제 성능을 판단하기 어렵습니다.")

# 4. AI 투자 규모 (차트)
s = blank(prs)
title(s, "AI 투자 규모 (2025년, 억 달러)")
cd = CategoryChartData()
cd.categories = ["글로벌 기업 AI", "미국 민간", "생성형 AI", "중국 민간"]
cd.add_series("투자액", (5817, 2859, 1709, 124))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.5), Inches(1.8),
                        Inches(10.33), Inches(4.5), cd)
gf.name = "Chart"
ch = gf.chart
ch.has_title = False
ch.has_legend = False
ch.font.size, ch.font.name = Pt(16), BODY_FONT
ch.font.color.rgb = BLACK
plot = ch.plots[0]
plot.gap_width = 80
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format, dl.number_format_is_linked = "#,##0", False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size, dl.font.bold = Pt(18), True
for i, pt in enumerate(plot.series[0].points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = ACCENTS[i]
va = ch.value_axis
va.has_major_gridlines = False
va.visible = False
ca = ch.category_axis
ca.format.line.color.rgb = BLACK
ca.format.line.width = Pt(1)
ca.tick_labels.font.size = Pt(16)
text(s, 1.5, 6.35, 7.5, 0.4, "출처: Stanford HAI, AI Index 2026", 14, align=PP_ALIGN.LEFT,
     name="ChartSource")
page_no(s, 4, TOTAL)
set_notes(s, "2025년 글로벌 기업 AI 투자는 5,817억 달러로 전년 대비 130% 늘었고, 생성형 AI "
             "투자는 약 5배인 1,709억 달러였습니다. 미국 민간 투자 2,859억 달러는 중국 "
             "124억 달러의 23배가 넘습니다. 미국은 데이터센터도 5,427개로 다른 나라의 10배 "
             "이상입니다. 막대들은 서로 포함 관계가 있으므로 합계가 아니라 규모 비교로 봐 주세요.")

# 5. 생산성·노동시장 (불릿 카드)
s = blank(prs)
title(s, "생산성·노동시장 영향")
box_x, box_y, box_w, box_h = 1.0, 1.9, SLIDE_W - 2.0, 4.55
card = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, box_x, box_y, box_w, box_h, fill=WHITE,
             line=BLACK, name="BulletCard")
card.adjustments[0] = 0.1 / box_h
shape(s, MSO_SHAPE.RECTANGLE, box_x + 0.15, box_y + 0.3, 0.12, box_h - 0.6, fill=CORAL,
      name="AccentBar")
text(s, box_x + 0.55, box_y + 0.3, box_w - 0.85, box_h - 0.6, [
    "• 생성형 AI 보급률 53%: PC·인터넷보다 빠른 확산",
    "• AI 노출 산업 생산성 +34% (비노출 산업 +24%)",
    "• AI 역량 보유자 임금 프리미엄 62%",
    "• 전체 고용 감소 효과는 0.4% 미만으로 제한적",
    "• 미국 22–25세 개발자 고용은 약 20% 감소",
    "• 일자리 긍정 전망: 전문가 73% vs 대중 23%",
], 20, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, name="Bullets", space_after=10)
page_no(s, 5, TOTAL)
set_notes(s, "생성형 AI는 3년 만에 보급률 53%에 도달해 PC나 인터넷보다 빠르게 퍼졌고, "
             "조사 대상 조직의 88%가 AI를 씁니다. PwC에 따르면 AI 노출이 큰 산업의 생산성은 "
             "2018년 대비 34% 늘어 노출이 적은 산업(24%)을 앞섰습니다. 전체 고용 감소는 "
             "0.4% 미만으로 작지만, 신입·청년층에 영향이 집중돼 미국 22–25세 개발자 고용은 "
             "2024년 대비 약 20% 줄었습니다. 전문가와 대중의 인식 차이도 큽니다.")

# 6. 규제 비교 (표)
s = blank(prs)
title(s, "AI 규제 비교: EU vs 한국")
rows = [
    ["구분", "EU AI Act", "한국 AI 기본법"],
    ["시행", "2026.8.2 전면 적용", "2026.1.22 시행"],
    ["대상", "범용 AI(GPAI) 제공자", "고영향·생성형 AI"],
    ["핵심 의무", "기술문서·학습데이터 요약", "AI 생성물 고지·영향평가"],
    ["집행", "AI Office 집행 권한", "과태료 최대 3천만 원"],
]
widths = [2.33, 4.5, 4.5]
tg = s.shapes.add_table(len(rows), 3, Inches(1.0), Inches(1.9), Inches(sum(widths)),
                        Inches(0.75 * len(rows)))
tg.name = "Table"
tbl = tg.table
tbl.first_row = False
tbl.horz_banding = False
for c, w in enumerate(widths):
    tbl.columns[c].width = Inches(w)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.75)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENTS[c] if r == 0 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if r == 0:
            style(p.add_run(), 22, True, WHITE, HEAD_FONT)
        else:
            style(p.add_run(), 18, c == 0, BLACK)
        p.runs[0].text = val
        table_borders(cell)
page_no(s, 6, TOTAL)
set_notes(s, "EU AI Act는 범용 AI 제공자 의무가 2025년 8월 2일부터, 대부분의 의무와 AI Office의 "
             "집행 권한은 2026년 8월 2일부터 적용됐습니다. 한국 AI 기본법은 2024년 12월 26일 "
             "통과돼 2026년 1월 22일 시행됐으며, 산업 진흥과 기본 의무를 함께 담고 과태료는 "
             "최대 3천만 원, 최소 1년의 계도 기간을 둡니다. 한편 보고된 AI 사고는 2024년 "
             "233건에서 2025년 362건으로 늘어 규제의 필요성이 커지고 있습니다.")

# 7. 핵심 이슈와 전망
s = blank(prs)
title(s, "핵심 이슈와 전망")
cards(s, [
    dict(head="실사용 검증", icon=MSO_SHAPE.DIAMOND, sub="점수 ≠ 현실",
         body=["벤치마크보다", "실제 환경 검증이 중요"]),
    dict(head="에이전트 확산", icon=MSO_SHAPE.GEAR_6, sub="79% 도입 난관",
         body=["병목은 거버넌스와", "변화 관리"]),
    dict(head="혜택 편중", icon=MSO_SHAPE.ISOSCELES_TRIANGLE, sub="신입층 압박",
         body=["생산성 이익은", "선도 기업에 집중"]),
    dict(head="물리적 한계", icon=MSO_SHAPE.LIGHTNING_BOLT, sub="전력·반도체",
         body=["2030 데이터센터", "전력 945TWh 전망"]),
])
page_no(s, 7, TOTAL)
set_notes(s, "첫째, 성능은 계속 빨라지지만 들쭉날쭉해 실제 환경 검증이 더 중요해졌습니다. "
             "둘째, 에이전트가 다음 도입 물결이지만 조직의 79%가 도입 어려움을 겪습니다. "
             "셋째, 생산성 이익은 선도 기업과 숙련 인력에 몰리고 신입층이 가장 큰 압박을 받습니다. "
             "넷째, 데이터센터 전력 수요는 2030년 약 945TWh로 두 배가 될 전망이고, 선도 AI 칩 "
             "대부분을 TSMC가 생산하는 공급망 집중도 위험 요인입니다. 안전과 정책이 성능을 "
             "따라가지 못하는 점도 함께 기억할 과제입니다.")

# 8. 마무리
s = blank(prs)
text(s, 1.0, 2.2, SLIDE_W - 2.0, 1.2, "감사합니다", 54, True, latin=HEAD_FONT,
     anchor=MSO_ANCHOR.MIDDLE, name="ClosingTitle")
text(s, 1.0, 3.45, SLIDE_W - 2.0, 0.7, "질의응답 · Q&A", 26, True, YELLOW, HEAD_FONT,
     anchor=MSO_ANCHOR.MIDDLE, name="ClosingSub")
dots(s, 4.5)
text(s, 1.0, 5.6, SLIDE_W - 2.0, 0.45, f"자료 출처 전체 목록: {SRC}", 14, name="ClosingSource")
set_notes(s, "정리하면, AI는 성능과 도입 모두 빠르게 커지고 있지만 검증, 거버넌스, 혜택의 "
             "분배, 에너지라는 과제가 함께 커지고 있습니다. 질문 받겠습니다.")


# ---------- 점검: 여백·겹침·텍스트 넘침 (ppt-rule-v1 §3) ----------

def text_height_in(tf, box_w_in):
    """줄바꿈을 추정해 텍스트가 차지할 높이(in)를 계산. 한글·전각 1em, 그 외 0.55em."""
    inner = box_w_in - (tf.margin_left + tf.margin_right) / 914400
    total = (tf.margin_top + tf.margin_bottom) / 914400
    for p in tf.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        em = size / 72
        width = sum(em * (1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.55)
                    for ch in p.text)
        lines = max(1, math.ceil(width / inner))
        total += lines * em * 1.2 + (p.space_after.pt / 72 if p.space_after else 0)
    return total


def check(prs):
    sw, sh = prs.slide_width / 914400, prs.slide_height / 914400
    problems = []
    for n, slide in enumerate(prs.slides, 1):
        boxes = []
        for shp in slide.shapes:
            x, y = shp.left / 914400, shp.top / 914400
            w, h = shp.width / 914400, shp.height / 914400
            if x < MARGIN - 1e-3 or y < MARGIN - 1e-3 or x + w > sw - MARGIN + 1e-3 \
                    or y + h > sh - MARGIN + 1e-3:
                problems.append(f"[{n}] {shp.name}: 여백 0.4in 미만")
            if shp.has_text_frame and shp.text_frame.text.strip():
                need = text_height_in(shp.text_frame, w)
                if need > h + 0.02:
                    problems.append(f"[{n}] {shp.name}: 텍스트 넘침 ({need:.2f} > {h:.2f} in)")
            boxes.append((shp.name, x, y, x + w, y + h))
        # 한쪽이 다른 쪽을 완전히 포함하면(배경 도형 위 텍스트) 의도한 겹침으로 본다
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                ix = min(a[3], b[3]) - max(a[1], b[1])
                iy = min(a[4], b[4]) - max(a[2], b[2])
                if ix > 0.01 and iy > 0.01:
                    a_in_b = a[1] >= b[1] and a[2] >= b[2] and a[3] <= b[3] and a[4] <= b[4]
                    b_in_a = b[1] >= a[1] and b[2] >= a[2] and b[3] <= a[3] and b[4] <= a[4]
                    if not (a_in_b or b_in_a):
                        problems.append(f"[{n}] 겹침: {a[0]} ↔ {b[0]}")
    return problems


issues = check(prs)
if issues:
    print("점검 실패:\n  " + "\n  ".join(issues))
    sys.exit(1)
print("점검 통과: 여백·겹침·텍스트 넘침 문제 없음")
print("저장:", save(prs, next_version_path("presentation", "AI_발전_발표")))
